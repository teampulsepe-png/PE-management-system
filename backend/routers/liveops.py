import io
import json
import re
import uuid
from datetime import datetime, timezone, timedelta
from typing import Optional

from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload

from backend.db.session import get_db
from backend.db.models import (
    LiveOpsBusinessUnit, LiveOpsTicketType, LiveOpsUseCase,
    LiveOpsApproverConfig, LiveOpsMemberRole, LiveOpsSlaConfig,
    LiveOpsTicket, LiveOpsAssignment, LiveOpsComment, LiveOpsTicketEvent,
    TeamMember, Notification,
)
from backend.schemas import (
    LiveOpsBusinessUnitOut, LiveOpsTicketTypeOut, LiveOpsUseCaseOut,
    LiveOpsApproverConfigOut, LiveOpsApproverConfigCreate,
    LiveOpsMemberRoleOut, LiveOpsMemberRoleCreate,
    LiveOpsSlaConfigOut, LiveOpsSlaConfigUpdate,
    LiveOpsTicketCreate, LiveOpsTicketListItem, LiveOpsTicketDetail,
    LiveOpsAssignmentOut, LiveOpsCommentOut, LiveOpsCommentCreate,
    LiveOpsEventOut, LiveOpsAssignableMember, LiveOpsAnalyticsSummary,
    LiveOpsApproveBody, LiveOpsRejectBody, LiveOpsAssignBody,
    LiveOpsReassignBody, LiveOpsHoldBody, LiveOpsCancelBody,
    LiveOpsCompleteBody, LiveOpsConfirmCompleteBody, LiveOpsRejectAssignmentBody, LiveOpsOverrideUrgencyBody,
    LiveOpsPushBackBody, LiveOpsTicketPatch,
)
from backend import sse, push_service

router = APIRouter(prefix="/api/v1/liveops", tags=["liveops"])

_MENTION_RE = re.compile(r'@\[([^\]]+)\]\(([^)]+)\)')


def _extract_mentions(body: str) -> list[tuple[str, str]]:
    """Returns [(display_name, member_id), ...] for every @[Name](id) in the text."""
    return _MENTION_RE.findall(body)


# ── helpers ───────────────────────────────────────────────────────────────────

def _now() -> datetime:
    return datetime.now(timezone.utc)


def _event(db: Session, ticket_id: str, actor_id: Optional[str], event_type: str, meta: dict = None):
    db.add(LiveOpsTicketEvent(
        id=str(uuid.uuid4()),
        ticket_id=ticket_id,
        actor_id=actor_id,
        event_type=event_type,
        event_data=json.dumps(meta) if meta else None,
    ))


def _system_comment(db: Session, ticket_id: str, body: str):
    db.add(LiveOpsComment(
        id=str(uuid.uuid4()),
        ticket_id=ticket_id,
        author_id=None,
        body=body,
        is_system=True,
    ))


def _sla_deadline(db: Session, urgency: str) -> Optional[datetime]:
    cfg = db.query(LiveOpsSlaConfig).filter_by(urgency=urgency).first()
    if not cfg:
        return None
    return _now() + timedelta(hours=cfg.response_hours)


def _next_ticket_number(db: Session) -> int:
    from sqlalchemy import func
    result = db.query(func.coalesce(func.max(LiveOpsTicket.ticket_number), 0)).scalar()
    return result + 1


def _load_ticket(db: Session, ticket_id: str) -> LiveOpsTicket:
    ticket = (
        db.query(LiveOpsTicket)
        .options(
            joinedload(LiveOpsTicket.ticket_type),
            joinedload(LiveOpsTicket.use_case),
            joinedload(LiveOpsTicket.business_unit),
            joinedload(LiveOpsTicket.submitted_by),
            joinedload(LiveOpsTicket.workstream_lead),
            joinedload(LiveOpsTicket.team_lead),
            joinedload(LiveOpsTicket.wl_actor),
            joinedload(LiveOpsTicket.tl_actor),
            joinedload(LiveOpsTicket.lo_actor),
            joinedload(LiveOpsTicket.currently_with),
            joinedload(LiveOpsTicket.cancelled_by),
            joinedload(LiveOpsTicket.assignments).joinedload(LiveOpsAssignment.assignee),
            joinedload(LiveOpsTicket.assignments).joinedload(LiveOpsAssignment.assigned_by),
            joinedload(LiveOpsTicket.comments).joinedload(LiveOpsComment.author),
            joinedload(LiveOpsTicket.events).joinedload(LiveOpsTicketEvent.actor),
        )
        .filter(LiveOpsTicket.id == ticket_id)
        .first()
    )
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")
    return ticket


def _get_member_roles(db: Session, member_id: str) -> set[str]:
    rows = db.query(LiveOpsMemberRole).filter_by(member_id=member_id).all()
    return {r.role for r in rows}


def _is_system_admin(member: TeamMember) -> bool:
    if getattr(member, 'has_admin_access', False):
        return True
    try:
        return member.role is not None and member.role.name == 'admin'
    except Exception:
        return False


def _lo_manager_id(db: Session) -> Optional[str]:
    """Returns the first LO Manager member ID, or None."""
    row = db.query(LiveOpsMemberRole).filter_by(role="lo_manager").first()
    return row.member_id if row else None


def _notify(
    db: Session,
    recipient_id: Optional[str],
    ntype: str,
    msg: str,
    actor_name: Optional[str] = None,
    entity_type: Optional[str] = None,
    entity_id: Optional[str] = None,
    notif_meta: Optional[dict] = None,
) -> Optional[tuple[str, str]]:
    """Persist a notification and return (member_id, sse_payload) for push after commit."""
    if not recipient_id:
        return None
    notif_id = str(uuid.uuid4())
    db.add(Notification(
        id=notif_id,
        recipient_member_id=recipient_id,
        type=ntype,
        entity_type=entity_type,
        entity_id=entity_id,
        notif_meta=notif_meta,
        triggered_by_name=actor_name,
        message=msg,
        is_read=False,
    ))
    payload = json.dumps({
        "id": notif_id, "type": ntype,
        "entity_type": entity_type, "entity_id": entity_id, "notif_meta": notif_meta,
        "triggered_by_name": actor_name,
        "message": msg, "is_read": False,
        "created_at": _now().isoformat(),
    })
    return (recipient_id, payload)


def _push_all(pushes: list[tuple[str, str]]) -> None:
    for member_id, payload in pushes:
        sse.push(member_id, payload)


# ── lookup tables ─────────────────────────────────────────────────────────────

@router.get("/business-units", response_model=list[LiveOpsBusinessUnitOut])
def list_business_units(db: Session = Depends(get_db)):
    return db.query(LiveOpsBusinessUnit).order_by(LiveOpsBusinessUnit.name).all()


@router.get("/ticket-types", response_model=list[LiveOpsTicketTypeOut])
def list_ticket_types(db: Session = Depends(get_db)):
    return db.query(LiveOpsTicketType).order_by(LiveOpsTicketType.sort_order, LiveOpsTicketType.name).all()


@router.get("/use-cases", response_model=list[LiveOpsUseCaseOut])
def list_use_cases(business_unit_id: Optional[str] = None, db: Session = Depends(get_db)):
    q = db.query(LiveOpsUseCase).options(joinedload(LiveOpsUseCase.business_unit))
    if business_unit_id:
        q = q.filter(LiveOpsUseCase.business_unit_id == business_unit_id)
    return q.order_by(LiveOpsUseCase.name).all()


# ── approver config ───────────────────────────────────────────────────────────

@router.get("/admin/approver-config", response_model=list[LiveOpsApproverConfigOut])
def list_approver_config(db: Session = Depends(get_db)):
    return (
        db.query(LiveOpsApproverConfig)
        .options(
            joinedload(LiveOpsApproverConfig.ticket_type),
            joinedload(LiveOpsApproverConfig.business_unit),
            joinedload(LiveOpsApproverConfig.workstream_lead),
            joinedload(LiveOpsApproverConfig.team_lead),
        )
        .all()
    )


@router.post("/admin/approver-config", response_model=LiveOpsApproverConfigOut, status_code=201)
def create_approver_config(body: LiveOpsApproverConfigCreate, db: Session = Depends(get_db)):
    existing = db.query(LiveOpsApproverConfig).filter_by(
        ticket_type_id=body.ticket_type_id,
        business_unit_id=body.business_unit_id,
    ).first()
    if existing:
        raise HTTPException(status_code=409, detail="Config already exists for this type + BU combination")
    cfg = LiveOpsApproverConfig(id=str(uuid.uuid4()), **body.model_dump())
    db.add(cfg)
    db.commit()
    return (
        db.query(LiveOpsApproverConfig)
        .options(
            joinedload(LiveOpsApproverConfig.ticket_type),
            joinedload(LiveOpsApproverConfig.business_unit),
            joinedload(LiveOpsApproverConfig.workstream_lead),
            joinedload(LiveOpsApproverConfig.team_lead),
        )
        .filter(LiveOpsApproverConfig.id == cfg.id)
        .first()
    )


@router.delete("/admin/approver-config/{config_id}", status_code=204)
def delete_approver_config(config_id: str, db: Session = Depends(get_db)):
    cfg = db.query(LiveOpsApproverConfig).filter_by(id=config_id).first()
    if not cfg:
        raise HTTPException(status_code=404, detail="Config not found")
    db.delete(cfg)
    db.commit()


# ── member roles ──────────────────────────────────────────────────────────────

@router.get("/admin/member-roles", response_model=list[LiveOpsMemberRoleOut])
def list_member_roles(db: Session = Depends(get_db)):
    return (
        db.query(LiveOpsMemberRole)
        .options(joinedload(LiveOpsMemberRole.member))
        .order_by(LiveOpsMemberRole.role)
        .all()
    )


@router.post("/admin/member-roles", response_model=LiveOpsMemberRoleOut, status_code=201)
def create_member_role(body: LiveOpsMemberRoleCreate, db: Session = Depends(get_db)):
    valid_roles = {"workstream_lead", "team_lead", "lo_manager", "platform_engineer", "liveops_engineer"}
    if body.role not in valid_roles:
        raise HTTPException(status_code=422, detail=f"role must be one of {sorted(valid_roles)}")
    existing = db.query(LiveOpsMemberRole).filter_by(member_id=body.member_id, role=body.role).first()
    if existing:
        raise HTTPException(status_code=409, detail="Member already has this role")
    row = LiveOpsMemberRole(id=str(uuid.uuid4()), **body.model_dump())
    db.add(row)
    db.commit()
    return (
        db.query(LiveOpsMemberRole)
        .options(joinedload(LiveOpsMemberRole.member))
        .filter(LiveOpsMemberRole.id == row.id)
        .first()
    )


@router.delete("/admin/member-roles/{role_id}", status_code=204)
def delete_member_role(role_id: str, db: Session = Depends(get_db)):
    row = db.query(LiveOpsMemberRole).filter_by(id=role_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Role assignment not found")
    db.delete(row)
    db.commit()


# ── SLA config ────────────────────────────────────────────────────────────────

@router.get("/admin/sla-config", response_model=list[LiveOpsSlaConfigOut])
def get_sla_config(db: Session = Depends(get_db)):
    return db.query(LiveOpsSlaConfig).order_by(LiveOpsSlaConfig.urgency).all()


@router.patch("/admin/sla-config/{urgency}", response_model=LiveOpsSlaConfigOut)
def update_sla_config(urgency: str, body: LiveOpsSlaConfigUpdate, db: Session = Depends(get_db)):
    cfg = db.query(LiveOpsSlaConfig).filter_by(urgency=urgency).first()
    if not cfg:
        raise HTTPException(status_code=404, detail="SLA config not found")
    cfg.response_hours = body.response_hours
    cfg.updated_at = _now()
    db.commit()
    db.refresh(cfg)
    return cfg


# ── assignable members ────────────────────────────────────────────────────────

@router.get("/assignable-members", response_model=list[LiveOpsAssignableMember])
def assignable_members(db: Session = Depends(get_db)):
    """All engineers (PE + LOE) with their active assignment counts — for the assignment UI."""
    eng_ids = {
        r.member_id
        for r in db.query(LiveOpsMemberRole)
        .filter(LiveOpsMemberRole.role.in_(["platform_engineer", "liveops_engineer"]))
        .all()
    }
    if not eng_ids:
        return []

    members = db.query(TeamMember).filter(TeamMember.id.in_(eng_ids)).all()

    active_counts: dict[str, int] = {}
    active_assignments = (
        db.query(LiveOpsAssignment)
        .filter(
            LiveOpsAssignment.assignee_id.in_(eng_ids),
            LiveOpsAssignment.status.in_(["pending", "active"]),
        )
        .all()
    )
    for a in active_assignments:
        active_counts[a.assignee_id] = active_counts.get(a.assignee_id, 0) + 1

    result = []
    for m in sorted(members, key=lambda x: x.name):
        result.append(LiveOpsAssignableMember(
            id=m.id,
            name=m.name,
            active_assignment_count=active_counts.get(m.id, 0),
        ))
    return result


# ── tickets — list ────────────────────────────────────────────────────────────

@router.get("/tickets", response_model=list[LiveOpsTicketListItem])
def list_tickets(
    status: Optional[str] = None,
    urgency: Optional[str] = None,
    ticket_type_id: Optional[str] = None,
    business_unit_id: Optional[str] = None,
    submitted_by_id: Optional[str] = None,
    assigned_to_id: Optional[str] = None,
    sla_breached: Optional[bool] = None,
    db: Session = Depends(get_db),
):
    q = (
        db.query(LiveOpsTicket)
        .options(
            joinedload(LiveOpsTicket.ticket_type),
            joinedload(LiveOpsTicket.use_case),
            joinedload(LiveOpsTicket.business_unit),
            joinedload(LiveOpsTicket.submitted_by),
            joinedload(LiveOpsTicket.currently_with),
            joinedload(LiveOpsTicket.workstream_lead),
            joinedload(LiveOpsTicket.team_lead),
        )
    )
    if status:
        q = q.filter(LiveOpsTicket.status == status)
    if urgency:
        q = q.filter(LiveOpsTicket.urgency == urgency)
    if ticket_type_id:
        q = q.filter(LiveOpsTicket.ticket_type_id == ticket_type_id)
    if business_unit_id:
        q = q.filter(LiveOpsTicket.business_unit_id == business_unit_id)
    if submitted_by_id:
        q = q.filter(LiveOpsTicket.submitted_by_id == submitted_by_id)
    if assigned_to_id:
        assigned_ticket_ids = {
            a.ticket_id
            for a in db.query(LiveOpsAssignment)
            .filter_by(assignee_id=assigned_to_id)
            .filter(LiveOpsAssignment.status.in_(["pending", "active"]))
            .all()
        }
        q = q.filter(LiveOpsTicket.id.in_(assigned_ticket_ids))
    if sla_breached is not None:
        now = _now()
        if sla_breached:
            q = q.filter(
                LiveOpsTicket.sla_deadline.isnot(None),
                LiveOpsTicket.sla_deadline < now,
                LiveOpsTicket.status.notin_(["completed", "rejected", "cancelled"]),
            )
        else:
            q = q.filter(
                LiveOpsTicket.sla_deadline.isnot(None),
                LiveOpsTicket.sla_deadline >= now,
            )

    return q.order_by(LiveOpsTicket.created_at.desc()).all()


# ── tickets — detail ──────────────────────────────────────────────────────────

@router.get("/tickets/{ticket_id}", response_model=LiveOpsTicketDetail)
def get_ticket(ticket_id: str, db: Session = Depends(get_db)):
    return _load_ticket(db, ticket_id)


# ── tickets — create ──────────────────────────────────────────────────────────

@router.post("/tickets", response_model=LiveOpsTicketDetail, status_code=201)
def create_ticket(body: LiveOpsTicketCreate, member_email: str, db: Session = Depends(get_db)):
    member = db.query(TeamMember).filter_by(email=member_email).first()
    if not member:
        raise HTTPException(status_code=404, detail="Member not found")

    number = _next_ticket_number(db)
    ticket = LiveOpsTicket(
        id=str(uuid.uuid4()),
        ticket_number=number,
        ticket_type_id=body.ticket_type_id,
        use_case_id=body.use_case_id,
        business_unit_id=body.business_unit_id,
        urgency=body.urgency,
        description=body.description,
        submitted_by_id=member.id,
        status="draft",
    )
    db.add(ticket)
    db.flush()

    pushes: list[tuple[str, str]] = []
    if body.submit:
        pushes = _do_submit(db, ticket, member)

    db.commit()
    _push_all(pushes)
    return _load_ticket(db, ticket.id)


# ── tickets — submit (draft → pending_wl) ─────────────────────────────────────

@router.post("/tickets/{ticket_id}/submit", response_model=LiveOpsTicketDetail)
def submit_ticket(ticket_id: str, member_email: str, db: Session = Depends(get_db)):
    member = db.query(TeamMember).filter_by(email=member_email).first()
    if not member:
        raise HTTPException(status_code=404, detail="Member not found")
    ticket = db.query(LiveOpsTicket).filter_by(id=ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")
    if ticket.submitted_by_id != member.id and not _is_system_admin(member):
        raise HTTPException(status_code=403, detail="Only the submitter can submit this ticket")
    if ticket.status != "draft":
        raise HTTPException(status_code=409, detail=f"Ticket is already {ticket.status}")

    pushes = _do_submit(db, ticket, member)
    db.commit()
    _push_all(pushes)
    return _load_ticket(db, ticket.id)


def _do_submit(db: Session, ticket: LiveOpsTicket, member: TeamMember) -> list[tuple[str, str]]:
    """Shared logic for first-time submission and re-submission after push back."""
    cfg = db.query(LiveOpsApproverConfig).filter_by(
        ticket_type_id=ticket.ticket_type_id,
        business_unit_id=ticket.business_unit_id,
    ).first()

    ticket.workstream_lead_id = cfg.workstream_lead_id if cfg else None
    ticket.team_lead_id       = cfg.team_lead_id if cfg else None

    # Reset per-stage audit for re-submission
    ticket.wl_status   = "pending"
    ticket.wl_actor_id = None
    ticket.wl_acted_at = None
    ticket.wl_reason   = None

    ticket.tl_status   = "na" if ticket.team_lead_id is None else "pending"
    ticket.tl_actor_id = None
    ticket.tl_acted_at = None
    ticket.tl_reason   = None

    ticket.lo_status   = "pending"
    ticket.lo_actor_id = None
    ticket.lo_acted_at = None
    ticket.lo_reason   = None

    ticket.submitted_at = _now()
    ticket.sla_deadline = _sla_deadline(db, ticket.urgency)

    pushes = []

    if ticket.workstream_lead_id:
        ticket.status = "pending_wl"
        ticket.currently_with_id = ticket.workstream_lead_id
        _event(db, ticket.id, member.id, "submitted", {
            "workstream_lead_id": ticket.workstream_lead_id,
            "team_lead_id": ticket.team_lead_id,
        })
        _system_comment(db, ticket.id, f"Ticket submitted by {member.name}.")
        n = _notify(db, ticket.workstream_lead_id, "liveops_ticket",
                    f"LO-{ticket.ticket_number:04d} needs your approval — submitted by {member.name}",
                    member.name, entity_type="liveops_ticket", entity_id=ticket.id)
        if n: pushes.append(n)
    else:
        # No WL assigned — skip approval chain and go straight to LO Manager
        ticket.wl_status = "na"
        ticket.tl_status = "na"
        ticket.status = "pending_lo"
        lo_id = _lo_manager_id(db)
        ticket.currently_with_id = lo_id
        _event(db, ticket.id, member.id, "submitted_direct_lo", {})
        _system_comment(db, ticket.id,
            f"Ticket submitted by {member.name}. No approver configured — forwarded directly to LO Manager.")
        n = _notify(db, lo_id, "liveops_ticket",
                    f"LO-{ticket.ticket_number:04d} submitted by {member.name} — no WL configured, needs your decision",
                    member.name, entity_type="liveops_ticket", entity_id=ticket.id)
        if n: pushes.append(n)

    return pushes


# ── tickets — patch (draft only) ─────────────────────────────────────────────

@router.patch("/tickets/{ticket_id}", response_model=LiveOpsTicketDetail)
def patch_ticket(ticket_id: str, body: LiveOpsTicketPatch, member_email: str, db: Session = Depends(get_db)):
    member = db.query(TeamMember).filter_by(email=member_email).first()
    if not member:
        raise HTTPException(status_code=404, detail="Member not found")
    ticket = db.query(LiveOpsTicket).filter_by(id=ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")
    if ticket.status != "draft":
        raise HTTPException(status_code=409, detail="Only draft tickets can be edited")
    if ticket.submitted_by_id != member.id and not _is_system_admin(member):
        raise HTTPException(status_code=403, detail="Not the ticket submitter")
    if body.description is not None:
        ticket.description = body.description
    if body.urgency is not None:
        if body.urgency not in ("low", "medium", "high"):
            raise HTTPException(status_code=422, detail="urgency must be low, medium, or high")
        ticket.urgency = body.urgency
    ticket.updated_at = _now()
    db.commit()
    return _load_ticket(db, ticket.id)


# ── tickets — approve ─────────────────────────────────────────────────────────

@router.post("/tickets/{ticket_id}/approve", response_model=LiveOpsTicketDetail)
def approve_ticket(ticket_id: str, body: LiveOpsApproveBody, member_email: str, db: Session = Depends(get_db)):
    member = db.query(TeamMember).filter_by(email=member_email).first()
    if not member:
        raise HTTPException(status_code=404, detail="Member not found")

    ticket = db.query(LiveOpsTicket).filter_by(id=ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")

    roles = _get_member_roles(db, member.id)
    is_admin = _is_system_admin(member)
    is_lo_or_admin = "lo_manager" in roles or is_admin
    now = _now()
    pushes: list[tuple[str, str]] = []

    if ticket.status == "pending_wl":
        if ticket.workstream_lead_id != member.id and not is_lo_or_admin:
            raise HTTPException(status_code=403, detail="Not the assigned Workstream Lead")
        ticket.wl_status   = "approved"
        ticket.wl_actor_id = member.id
        ticket.wl_acted_at = now
        ticket.wl_reason   = body.reason

        if ticket.tl_status == "na":
            lo_id = _lo_manager_id(db)
            ticket.status = "pending_lo"
            ticket.currently_with_id = lo_id
            _event(db, ticket.id, member.id, "wl_approved", {"reason": body.reason, "next": "pending_lo"})
            _system_comment(db, ticket.id, f"Approved by Workstream Lead {member.name}. Forwarded to LO Manager.")
            n = _notify(db, lo_id, "liveops_ticket",
                        f"LO-{ticket.ticket_number:04d} approved by WL and awaiting your review", member.name,
                        entity_type="liveops_ticket", entity_id=ticket_id)
            if n: pushes.append(n)
        else:
            ticket.status = "pending_tl"
            ticket.currently_with_id = ticket.team_lead_id
            _event(db, ticket.id, member.id, "wl_approved", {"reason": body.reason, "next": "pending_tl"})
            _system_comment(db, ticket.id, f"Approved by Workstream Lead {member.name}. Forwarded to Team Lead.")
            n = _notify(db, ticket.team_lead_id, "liveops_ticket",
                        f"LO-{ticket.ticket_number:04d} needs your approval — WL approved", member.name,
                        entity_type="liveops_ticket", entity_id=ticket_id)
            if n: pushes.append(n)

    elif ticket.status == "pending_tl":
        if ticket.team_lead_id != member.id and not is_lo_or_admin:
            raise HTTPException(status_code=403, detail="Not the assigned Team Lead")
        lo_id = _lo_manager_id(db)
        ticket.tl_status   = "approved"
        ticket.tl_actor_id = member.id
        ticket.tl_acted_at = now
        ticket.tl_reason   = body.reason
        ticket.status = "pending_lo"
        ticket.currently_with_id = lo_id
        _event(db, ticket.id, member.id, "tl_approved", {"reason": body.reason})
        _system_comment(db, ticket.id, f"Approved by Team Lead {member.name}. Forwarded to LO Manager.")
        n = _notify(db, lo_id, "liveops_ticket",
                    f"LO-{ticket.ticket_number:04d} fully approved and awaiting your decision", member.name,
                    entity_type="liveops_ticket", entity_id=ticket_id)
        if n: pushes.append(n)

    elif ticket.status == "pending_lo":
        if not is_lo_or_admin:
            raise HTTPException(status_code=403, detail="Only LO Manager can approve at this stage")
        if not body.assignee_ids:
            raise HTTPException(status_code=422, detail="assignee_ids is required when approving at LO stage")

        assignees = []
        for aid in body.assignee_ids:
            assignee = db.query(TeamMember).filter_by(id=aid).first()
            if not assignee:
                raise HTTPException(status_code=404, detail=f"Assignee {aid} not found")
            role_ok = db.query(LiveOpsMemberRole).filter(
                LiveOpsMemberRole.member_id == aid,
                LiveOpsMemberRole.role.in_(["platform_engineer", "liveops_engineer"]),
            ).first()
            if not role_ok:
                raise HTTPException(status_code=422, detail=f"{assignee.name} must be a Platform Engineer or LiveOps Engineer")
            assignees.append(assignee)

        ticket.lo_status   = "approved"
        ticket.lo_actor_id = member.id
        ticket.lo_acted_at = now
        ticket.lo_reason   = body.reason
        ticket.status      = "open"
        ticket.currently_with_id = body.assignee_ids[0]

        assignee_names = []
        for assignee in assignees:
            db.add(LiveOpsAssignment(
                id=str(uuid.uuid4()),
                ticket_id=ticket_id,
                assignee_id=assignee.id,
                assigned_by_id=member.id,
                lo_message=body.lo_message,
                status="pending",
            ))
            assignee_names.append(assignee.name)
            msg = f"LO-{ticket.ticket_number:04d} has been assigned to you"
            if body.lo_message:
                msg += f" — {body.lo_message}"
            n = _notify(db, assignee.id, "liveops_ticket", msg, member.name,
                        entity_type="liveops_ticket", entity_id=ticket_id)
            if n: pushes.append(n)

        names_str = ", ".join(assignee_names)
        _event(db, ticket.id, member.id, "lo_approved_and_assigned", {
            "reason": body.reason,
            "assignee_ids": body.assignee_ids,
            "assignee_names": assignee_names,
        })
        _system_comment(db, ticket.id, f"Approved by LO Manager {member.name}. Assigned to {names_str}.")

    else:
        raise HTTPException(status_code=409, detail=f"Ticket cannot be approved in status '{ticket.status}'")

    ticket.updated_at = now
    db.commit()
    _push_all(pushes)
    return _load_ticket(db, ticket.id)


# ── tickets — push back ───────────────────────────────────────────────────────

@router.post("/tickets/{ticket_id}/push-back", response_model=LiveOpsTicketDetail)
def push_back_ticket(ticket_id: str, body: LiveOpsPushBackBody, member_email: str, db: Session = Depends(get_db)):
    member = db.query(TeamMember).filter_by(email=member_email).first()
    if not member:
        raise HTTPException(status_code=404, detail="Member not found")

    ticket = db.query(LiveOpsTicket).filter_by(id=ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")

    is_admin = _is_system_admin(member)
    now = _now()
    pushes: list[tuple[str, str]] = []

    if ticket.status == "pending_wl":
        if ticket.workstream_lead_id != member.id and not is_admin:
            raise HTTPException(status_code=403, detail="Not the assigned Workstream Lead")
        ticket.wl_status   = "pushed_back"
        ticket.wl_actor_id = member.id
        ticket.wl_acted_at = now
        ticket.wl_reason   = body.reason
        ticket.status      = "draft"
        ticket.currently_with_id = ticket.submitted_by_id
        _event(db, ticket.id, member.id, "wl_pushed_back", {"reason": body.reason})
        _system_comment(db, ticket.id,
            f"Pushed back to submitter by Workstream Lead {member.name}. Reason: {body.reason}")
        n = _notify(db, ticket.submitted_by_id, "liveops_ticket",
                    f"LO-{ticket.ticket_number:04d} was pushed back: {body.reason}", member.name,
                    entity_type="liveops_ticket", entity_id=ticket_id)
        if n: pushes.append(n)

    elif ticket.status == "pending_tl":
        if ticket.team_lead_id != member.id and not is_admin:
            raise HTTPException(status_code=403, detail="Not the assigned Team Lead")
        ticket.tl_status   = "pushed_back"
        ticket.tl_actor_id = member.id
        ticket.tl_acted_at = now
        ticket.tl_reason   = body.reason
        ticket.wl_status   = "pending"
        ticket.wl_actor_id = None
        ticket.wl_acted_at = None
        ticket.wl_reason   = None
        ticket.status      = "pending_wl"
        ticket.currently_with_id = ticket.workstream_lead_id
        _event(db, ticket.id, member.id, "tl_pushed_back", {"reason": body.reason})
        _system_comment(db, ticket.id,
            f"Pushed back to Workstream Lead by Team Lead {member.name}. Reason: {body.reason}")
        n = _notify(db, ticket.workstream_lead_id, "liveops_ticket",
                    f"LO-{ticket.ticket_number:04d} pushed back by Team Lead: {body.reason}", member.name,
                    entity_type="liveops_ticket", entity_id=ticket_id)
        if n: pushes.append(n)

    else:
        raise HTTPException(status_code=409, detail=f"Cannot push back ticket in status '{ticket.status}'")

    ticket.updated_at = now
    db.commit()
    _push_all(pushes)
    return _load_ticket(db, ticket.id)


# ── tickets — reject ──────────────────────────────────────────────────────────

@router.post("/tickets/{ticket_id}/reject", response_model=LiveOpsTicketDetail)
def reject_ticket(ticket_id: str, body: LiveOpsRejectBody, member_email: str, db: Session = Depends(get_db)):
    member = db.query(TeamMember).filter_by(email=member_email).first()
    if not member:
        raise HTTPException(status_code=404, detail="Member not found")

    ticket = db.query(LiveOpsTicket).filter_by(id=ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")

    roles = _get_member_roles(db, member.id)
    is_admin = _is_system_admin(member)
    is_lo_or_admin = "lo_manager" in roles or is_admin
    now = _now()
    pushes: list[tuple[str, str]] = []

    if ticket.status == "pending_wl":
        if ticket.workstream_lead_id != member.id and not is_lo_or_admin:
            raise HTTPException(status_code=403, detail="Not the assigned Workstream Lead")
        ticket.wl_status   = "rejected"
        ticket.wl_actor_id = member.id
        ticket.wl_acted_at = now
        ticket.wl_reason   = body.reason
        _event(db, ticket.id, member.id, "wl_rejected", {"reason": body.reason})
        _system_comment(db, ticket.id, f"Rejected by Workstream Lead {member.name}. Reason: {body.reason}")

    elif ticket.status == "pending_tl":
        if ticket.team_lead_id != member.id and not is_lo_or_admin:
            raise HTTPException(status_code=403, detail="Not the assigned Team Lead")
        ticket.tl_status   = "rejected"
        ticket.tl_actor_id = member.id
        ticket.tl_acted_at = now
        ticket.tl_reason   = body.reason
        _event(db, ticket.id, member.id, "tl_rejected", {"reason": body.reason})
        _system_comment(db, ticket.id, f"Rejected by Team Lead {member.name}. Reason: {body.reason}")

    elif ticket.status == "pending_lo":
        if not is_lo_or_admin:
            raise HTTPException(status_code=403, detail="Only LO Manager can reject at this stage")
        ticket.lo_status   = "rejected"
        ticket.lo_actor_id = member.id
        ticket.lo_acted_at = now
        ticket.lo_reason   = body.reason
        _event(db, ticket.id, member.id, "lo_rejected", {"reason": body.reason})
        _system_comment(db, ticket.id, f"Rejected by LO Manager {member.name}. Reason: {body.reason}")

    else:
        raise HTTPException(status_code=409, detail=f"Ticket cannot be rejected in status '{ticket.status}'")

    ticket.status = "rejected"
    ticket.currently_with_id = None
    ticket.updated_at = now
    n = _notify(db, ticket.submitted_by_id, "liveops_ticket",
                f"LO-{ticket.ticket_number:04d} was rejected: {body.reason}", member.name,
                entity_type="liveops_ticket", entity_id=ticket_id)
    if n: pushes.append(n)
    db.commit()
    _push_all(pushes)
    return _load_ticket(db, ticket.id)


# ── tickets — cancel ──────────────────────────────────────────────────────────

@router.post("/tickets/{ticket_id}/cancel", response_model=LiveOpsTicketDetail)
def cancel_ticket(ticket_id: str, body: LiveOpsCancelBody, member_email: str, db: Session = Depends(get_db)):
    member = db.query(TeamMember).filter_by(email=member_email).first()
    if not member:
        raise HTTPException(status_code=404, detail="Member not found")

    ticket = db.query(LiveOpsTicket).filter_by(id=ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")

    roles = _get_member_roles(db, member.id)
    is_submitter = ticket.submitted_by_id == member.id
    is_lo_or_admin = "lo_manager" in roles or _is_system_admin(member)

    if not is_submitter and not is_lo_or_admin:
        raise HTTPException(status_code=403, detail="Only the submitter or LO Manager can cancel")
    if ticket.status in ("completed", "rejected", "cancelled"):
        raise HTTPException(status_code=409, detail=f"Ticket is already {ticket.status}")

    now = _now()
    ticket.status = "cancelled"
    ticket.cancelled_by_id = member.id
    ticket.cancelled_at = now
    ticket.cancel_reason = body.reason
    ticket.currently_with_id = None
    ticket.updated_at = now

    _event(db, ticket.id, member.id, "cancelled", {"reason": body.reason})
    _system_comment(db, ticket.id, f"Cancelled by {member.name}. Reason: {body.reason}")
    db.commit()
    return _load_ticket(db, ticket.id)


# ── tickets — hold / resume ───────────────────────────────────────────────────

@router.post("/tickets/{ticket_id}/hold", response_model=LiveOpsTicketDetail)
def hold_ticket(ticket_id: str, body: LiveOpsHoldBody, member_email: str, db: Session = Depends(get_db)):
    member = db.query(TeamMember).filter_by(email=member_email).first()
    if not member:
        raise HTTPException(status_code=404, detail="Member not found")

    roles = _get_member_roles(db, member.id)
    if "lo_manager" not in roles and not _is_system_admin(member):
        raise HTTPException(status_code=403, detail="Only LO Manager can put tickets on hold")

    ticket = db.query(LiveOpsTicket).filter_by(id=ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")
    if ticket.status not in ("open", "in_progress", "pending_lo", "pending_review"):
        raise HTTPException(status_code=409, detail=f"Cannot put ticket on hold from status '{ticket.status}'")

    now = _now()
    ticket.status = "on_hold"
    ticket.on_hold_reason = body.reason
    ticket.put_on_hold_at = now
    ticket.updated_at = now

    _event(db, ticket.id, member.id, "put_on_hold", {"reason": body.reason})
    _system_comment(db, ticket.id, f"Put on hold by {member.name}. Reason: {body.reason}")
    db.commit()
    return _load_ticket(db, ticket.id)


@router.post("/tickets/{ticket_id}/resume", response_model=LiveOpsTicketDetail)
def resume_ticket(ticket_id: str, member_email: str, db: Session = Depends(get_db)):
    member = db.query(TeamMember).filter_by(email=member_email).first()
    if not member:
        raise HTTPException(status_code=404, detail="Member not found")

    roles = _get_member_roles(db, member.id)
    if "lo_manager" not in roles and not _is_system_admin(member):
        raise HTTPException(status_code=403, detail="Only LO Manager can resume tickets")

    ticket = db.query(LiveOpsTicket).filter_by(id=ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")
    if ticket.status != "on_hold":
        raise HTTPException(status_code=409, detail="Ticket is not on hold")

    now = _now()
    active = db.query(LiveOpsAssignment).filter_by(
        ticket_id=ticket.id, status="active"
    ).first()
    ticket.status = "in_progress" if active else "open"
    ticket.on_hold_reason = None
    ticket.put_on_hold_at = None
    ticket.updated_at = now

    _event(db, ticket.id, member.id, "resumed", {})
    _system_comment(db, ticket.id, f"Resumed by {member.name}.")
    db.commit()
    return _load_ticket(db, ticket.id)


# ── tickets — assign ──────────────────────────────────────────────────────────

@router.post("/tickets/{ticket_id}/assign", response_model=LiveOpsTicketDetail)
def assign_ticket(ticket_id: str, body: LiveOpsAssignBody, member_email: str, db: Session = Depends(get_db)):
    member = db.query(TeamMember).filter_by(email=member_email).first()
    if not member:
        raise HTTPException(status_code=404, detail="Member not found")

    roles = _get_member_roles(db, member.id)
    if "lo_manager" not in roles and not _is_system_admin(member):
        raise HTTPException(status_code=403, detail="Only LO Manager can assign tickets")

    ticket = db.query(LiveOpsTicket).filter_by(id=ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")
    if ticket.status not in ("open", "in_progress"):
        raise HTTPException(status_code=409, detail=f"Cannot assign ticket in status '{ticket.status}'")

    if not body.assignee_ids:
        raise HTTPException(status_code=422, detail="assignee_ids must not be empty")

    assignees = []
    for aid in body.assignee_ids:
        assignee = db.query(TeamMember).filter_by(id=aid).first()
        if not assignee:
            raise HTTPException(status_code=404, detail=f"Assignee {aid} not found")
        role_ok = db.query(LiveOpsMemberRole).filter(
            LiveOpsMemberRole.member_id == aid,
            LiveOpsMemberRole.role.in_(["platform_engineer", "liveops_engineer"]),
        ).first()
        if not role_ok:
            raise HTTPException(status_code=422, detail=f"{assignee.name} must be a Platform Engineer or LiveOps Engineer")
        already = db.query(LiveOpsAssignment).filter_by(ticket_id=ticket_id, assignee_id=aid).filter(
            LiveOpsAssignment.status.in_(["pending", "active"])
        ).first()
        if already:
            raise HTTPException(status_code=409, detail=f"{assignee.name} is already actively assigned to this ticket")
        assignees.append(assignee)

    now = _now()
    assignee_names = []
    for assignee in assignees:
        db.add(LiveOpsAssignment(
            id=str(uuid.uuid4()),
            ticket_id=ticket_id,
            assignee_id=assignee.id,
            assigned_by_id=member.id,
            lo_message=body.lo_message,
            status="pending",
        ))
        assignee_names.append(assignee.name)

    ticket.status = "in_progress"
    ticket.currently_with_id = body.assignee_ids[0]
    ticket.updated_at = now

    names_str = ", ".join(assignee_names)
    _event(db, ticket.id, member.id, "assigned", {
        "assignee_ids": body.assignee_ids, "assignee_names": assignee_names, "lo_message": body.lo_message,
    })
    _system_comment(db, ticket.id, f"Assigned to {names_str} by {member.name}.")
    pushes = []
    for assignee in assignees:
        msg = f"LO-{ticket.ticket_number:04d} has been assigned to you"
        if body.lo_message:
            msg += f" — {body.lo_message}"
        n = _notify(db, assignee.id, "liveops_ticket", msg, member.name,
                    entity_type="liveops_ticket", entity_id=ticket_id)
        if n: pushes.append(n)
    db.commit()
    _push_all(pushes)
    return _load_ticket(db, ticket.id)


# ── tickets — reassign ────────────────────────────────────────────────────────

@router.post("/tickets/{ticket_id}/reassign", response_model=LiveOpsTicketDetail)
def reassign_ticket(ticket_id: str, body: LiveOpsReassignBody, member_email: str, db: Session = Depends(get_db)):
    member = db.query(TeamMember).filter_by(email=member_email).first()
    if not member:
        raise HTTPException(status_code=404, detail="Member not found")

    roles = _get_member_roles(db, member.id)
    if "lo_manager" not in roles and not _is_system_admin(member):
        raise HTTPException(status_code=403, detail="Only LO Manager can reassign")

    ticket = db.query(LiveOpsTicket).filter_by(id=ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")

    new_assignee = db.query(TeamMember).filter_by(id=body.assignee_id).first()
    if not new_assignee:
        raise HTTPException(status_code=404, detail="New assignee not found")

    now = _now()
    old_assignments = db.query(LiveOpsAssignment).filter_by(ticket_id=ticket_id).filter(
        LiveOpsAssignment.status.in_(["pending", "active"])
    ).all()
    for a in old_assignments:
        a.status = "rejected"
        a.rejected_at = now
        a.rejection_reason = f"Reassigned to {new_assignee.name}" + (f": {body.reason}" if body.reason else "")

    new_assignment = LiveOpsAssignment(
        id=str(uuid.uuid4()),
        ticket_id=ticket_id,
        assignee_id=body.assignee_id,
        assigned_by_id=member.id,
        lo_message=body.lo_message,
        status="pending",
    )
    db.add(new_assignment)

    ticket.currently_with_id = body.assignee_id
    ticket.status = "in_progress"
    ticket.updated_at = now

    _event(db, ticket.id, member.id, "reassigned", {
        "new_assignee_id": body.assignee_id,
        "new_assignee_name": new_assignee.name,
        "reason": body.reason,
    })
    _system_comment(db, ticket.id, f"Reassigned to {new_assignee.name} by {member.name}.")
    db.commit()
    return _load_ticket(db, ticket.id)


# ── tickets — pick up ─────────────────────────────────────────────────────────

@router.post("/tickets/{ticket_id}/pickup", response_model=LiveOpsTicketDetail)
def pickup_ticket(ticket_id: str, member_email: str, db: Session = Depends(get_db)):
    member = db.query(TeamMember).filter_by(email=member_email).first()
    if not member:
        raise HTTPException(status_code=404, detail="Member not found")

    assignment = db.query(LiveOpsAssignment).filter_by(
        ticket_id=ticket_id,
        assignee_id=member.id,
        status="pending",
    ).first()
    if not assignment:
        raise HTTPException(status_code=404, detail="No pending assignment found for this member on this ticket")

    now = _now()
    assignment.status = "active"
    assignment.picked_up_at = now

    ticket = db.query(LiveOpsTicket).filter_by(id=ticket_id).first()
    ticket.status = "in_progress"
    ticket.currently_with_id = member.id
    ticket.updated_at = now

    _event(db, ticket_id, member.id, "picked_up", {})
    _system_comment(db, ticket_id, f"Picked up by {member.name}.")
    lo_id = _lo_manager_id(db)
    pushes = []
    n = _notify(db, lo_id, "liveops_ticket",
                f"LO-{ticket.ticket_number:04d} picked up by {member.name}", member.name,
                entity_type="liveops_ticket", entity_id=ticket_id)
    if n: pushes.append(n)
    db.commit()
    _push_all(pushes)
    return _load_ticket(db, ticket_id)


# ── tickets — complete ────────────────────────────────────────────────────────

@router.post("/tickets/{ticket_id}/complete", response_model=LiveOpsTicketDetail)
def complete_ticket(ticket_id: str, body: LiveOpsCompleteBody, member_email: str, db: Session = Depends(get_db)):
    member = db.query(TeamMember).filter_by(email=member_email).first()
    if not member:
        raise HTTPException(status_code=404, detail="Member not found")

    assignment = db.query(LiveOpsAssignment).filter_by(
        ticket_id=ticket_id,
        assignee_id=member.id,
        status="active",
    ).first()
    if not assignment:
        raise HTTPException(status_code=404, detail="No active assignment found for this member on this ticket")

    now = _now()
    assignment.status = "completed"
    assignment.completed_at = now

    ticket = db.query(LiveOpsTicket).filter_by(id=ticket_id).first()

    other_open = db.query(LiveOpsAssignment).filter_by(ticket_id=ticket_id).filter(
        LiveOpsAssignment.status.in_(["pending", "active"]),
        LiveOpsAssignment.id != assignment.id,
    ).count()

    if other_open == 0:
        ticket.status = "pending_review"
        ticket.currently_with_id = None
    ticket.updated_at = now

    note_text = f" Note: {body.note}" if body.note else ""
    _event(db, ticket_id, member.id, "work_completed", {"note": body.note})
    _system_comment(db, ticket_id, f"Work completed by {member.name}.{note_text} Pending LO review.")
    lo_id = _lo_manager_id(db)
    pushes = []
    n = _notify(db, lo_id, "liveops_ticket",
                f"LO-{ticket.ticket_number:04d} ready for review — work done by {member.name}", member.name,
                entity_type="liveops_ticket", entity_id=ticket_id)
    if n: pushes.append(n)
    db.commit()
    _push_all(pushes)
    return _load_ticket(db, ticket_id)


# ── tickets — confirm complete (LO Manager) ───────────────────────────────────

@router.post("/tickets/{ticket_id}/confirm-complete", response_model=LiveOpsTicketDetail)
def confirm_complete(ticket_id: str, body: LiveOpsConfirmCompleteBody, member_email: str, db: Session = Depends(get_db)):
    member = db.query(TeamMember).filter_by(email=member_email).first()
    if not member:
        raise HTTPException(status_code=404, detail="Member not found")

    roles = _get_member_roles(db, member.id)
    if "lo_manager" not in roles and not _is_system_admin(member):
        raise HTTPException(status_code=403, detail="Only LO Manager can confirm completion")

    ticket = db.query(LiveOpsTicket).filter_by(id=ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")
    if ticket.status != "pending_review":
        raise HTTPException(status_code=409, detail="Ticket is not pending review")

    now = _now()
    ticket.status = "completed"
    ticket.completed_at = now
    ticket.currently_with_id = None
    ticket.updated_at = now

    note_text = f" Note: {body.note}" if body.note else ""
    _event(db, ticket_id, member.id, "confirmed_complete", {"note": body.note})
    _system_comment(db, ticket_id, f"Completion confirmed by {member.name}.{note_text}")

    pushes = []
    n = _notify(db, ticket.submitted_by_id, "liveops_ticket",
                f"LO-{ticket.ticket_number:04d} has been completed and confirmed", member.name,
                entity_type="liveops_ticket", entity_id=ticket_id)
    if n: pushes.append(n)
    db.commit()
    _push_all(pushes)
    return _load_ticket(db, ticket_id)


# ── tickets — reject assignment ───────────────────────────────────────────────

@router.post("/tickets/{ticket_id}/reject-assignment", response_model=LiveOpsTicketDetail)
def reject_assignment(ticket_id: str, body: LiveOpsRejectAssignmentBody, member_email: str, db: Session = Depends(get_db)):
    member = db.query(TeamMember).filter_by(email=member_email).first()
    if not member:
        raise HTTPException(status_code=404, detail="Member not found")

    assignment = db.query(LiveOpsAssignment).filter_by(
        ticket_id=ticket_id,
        assignee_id=member.id,
    ).filter(LiveOpsAssignment.status.in_(["pending", "active"])).first()
    if not assignment:
        raise HTTPException(status_code=404, detail="No open assignment found for this member on this ticket")

    now = _now()
    assignment.status = "rejected"
    assignment.rejected_at = now
    assignment.rejection_reason = body.reason

    ticket = db.query(LiveOpsTicket).filter_by(id=ticket_id).first()

    other_open = db.query(LiveOpsAssignment).filter_by(ticket_id=ticket_id).filter(
        LiveOpsAssignment.status.in_(["pending", "active"]),
        LiveOpsAssignment.id != assignment.id,
    ).count()

    if other_open == 0:
        ticket.status = "open"
        ticket.currently_with_id = _lo_manager_id(db)

    ticket.updated_at = now
    _event(db, ticket_id, member.id, "assignment_rejected", {"reason": body.reason})
    _system_comment(db, ticket_id, f"Assignment declined by {member.name}. Reason: {body.reason}")
    lo_id = _lo_manager_id(db)
    pushes = []
    n = _notify(db, lo_id, "liveops_ticket",
                f"LO-{ticket.ticket_number:04d} assignment declined by {member.name}: {body.reason}", member.name,
                entity_type="liveops_ticket", entity_id=ticket_id)
    if n: pushes.append(n)
    db.commit()
    _push_all(pushes)
    return _load_ticket(db, ticket_id)


# ── tickets — override urgency ────────────────────────────────────────────────

@router.post("/tickets/{ticket_id}/override-urgency", response_model=LiveOpsTicketDetail)
def override_urgency(ticket_id: str, body: LiveOpsOverrideUrgencyBody, member_email: str, db: Session = Depends(get_db)):
    member = db.query(TeamMember).filter_by(email=member_email).first()
    if not member:
        raise HTTPException(status_code=404, detail="Member not found")

    roles = _get_member_roles(db, member.id)
    if "lo_manager" not in roles and not _is_system_admin(member):
        raise HTTPException(status_code=403, detail="Only LO Manager can override urgency")
    if body.urgency not in ("low", "medium", "high"):
        raise HTTPException(status_code=422, detail="urgency must be low, medium, or high")

    ticket = db.query(LiveOpsTicket).filter_by(id=ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")
    if ticket.status in ("completed", "rejected", "cancelled"):
        raise HTTPException(status_code=409, detail="Cannot change urgency on a closed ticket")

    old_urgency = ticket.urgency
    ticket.urgency = body.urgency
    ticket.sla_deadline = _sla_deadline(db, body.urgency)
    ticket.updated_at = _now()

    _event(db, ticket_id, member.id, "urgency_changed", {
        "from": old_urgency, "to": body.urgency, "reason": body.reason
    })
    _system_comment(db, ticket_id,
        f"Urgency changed from {old_urgency} to {body.urgency} by {member.name}."
        + (f" Reason: {body.reason}" if body.reason else "")
    )
    db.commit()
    return _load_ticket(db, ticket_id)


# ── comments ──────────────────────────────────────────────────────────────────

@router.get("/tickets/{ticket_id}/comments", response_model=list[LiveOpsCommentOut])
def list_comments(ticket_id: str, db: Session = Depends(get_db)):
    return (
        db.query(LiveOpsComment)
        .options(joinedload(LiveOpsComment.author))
        .filter_by(ticket_id=ticket_id)
        .order_by(LiveOpsComment.created_at)
        .all()
    )


@router.post("/tickets/{ticket_id}/comments", response_model=LiveOpsCommentOut, status_code=201)
def add_comment(
    ticket_id: str,
    body: LiveOpsCommentCreate,
    member_email: str,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
):
    member = db.query(TeamMember).filter_by(email=member_email).first()
    if not member:
        raise HTTPException(status_code=404, detail="Member not found")

    ticket = db.query(LiveOpsTicket).filter_by(id=ticket_id).first()
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")

    comment = LiveOpsComment(
        id=str(uuid.uuid4()),
        ticket_id=ticket_id,
        author_id=member.id,
        body=body.body,
        is_system=False,
    )
    db.add(comment)
    _event(db, ticket_id, member.id, "comment_added", {})

    # Parse @mentions → notify mentioned members
    sse_pushes: list[tuple[str, str]] = []
    mention_emails: list[str] = []
    for display_name, mentioned_id in _extract_mentions(body.body):
        if mentioned_id == member.id:
            continue
        mentioned = db.query(TeamMember).filter_by(id=mentioned_id).first()
        if not mentioned:
            continue
        if mentioned.email:
            mention_emails.append(mentioned.email)
        msg = f"{member.name} mentioned you in a comment on LO-{ticket.ticket_number:04d}"
        n = _notify(db, mentioned.id, "mention", msg, member.name,
                    entity_type="liveops_ticket", entity_id=ticket_id)
        if n:
            sse_pushes.append(n)

    db.commit()

    for mid, payload in sse_pushes:
        sse.push(mid, payload)

    if mention_emails:
        background_tasks.add_task(
            push_service.send_to_emails_sync,
            mention_emails,
            f"💬 LO-{ticket.ticket_number:04d}",
            f"{member.name} mentioned you in a comment",
            "/liveops",
        )

    return (
        db.query(LiveOpsComment)
        .options(joinedload(LiveOpsComment.author))
        .filter(LiveOpsComment.id == comment.id)
        .first()
    )


# ── reports export ────────────────────────────────────────────────────────────

@router.get("/reports/export")
def export_report(
    from_date: str = Query(..., description="YYYY-MM-DD"),
    to_date: str   = Query(..., description="YYYY-MM-DD"),
    format: str    = Query("excel", description="excel or pdf"),
    db: Session    = Depends(get_db),
):
    try:
        from_dt = datetime.strptime(from_date, "%Y-%m-%d").replace(tzinfo=timezone.utc)
        to_dt   = datetime.strptime(to_date,   "%Y-%m-%d").replace(hour=23, minute=59, second=59, tzinfo=timezone.utc)
    except ValueError:
        raise HTTPException(status_code=422, detail="Invalid date format — use YYYY-MM-DD")

    tickets = (
        db.query(LiveOpsTicket)
        .options(
            joinedload(LiveOpsTicket.ticket_type),
            joinedload(LiveOpsTicket.business_unit),
            joinedload(LiveOpsTicket.use_case),
            joinedload(LiveOpsTicket.submitted_by),
            joinedload(LiveOpsTicket.workstream_lead),
            joinedload(LiveOpsTicket.team_lead),
            joinedload(LiveOpsTicket.currently_with),
        )
        .filter(LiveOpsTicket.submitted_at >= from_dt, LiveOpsTicket.submitted_at <= to_dt)
        .order_by(LiveOpsTicket.submitted_at)
        .all()
    )

    headers_row = [
        "Ticket #", "Status", "Urgency", "Type", "Business Unit", "Use Case",
        "Submitted By", "Submitted At", "Workstream Lead", "Practice Lead",
        "Assigned To", "SLA Deadline", "Completed At",
    ]

    def _row(t: LiveOpsTicket) -> list:
        return [
            t.ticket_number,
            t.status.replace("_", " ").title(),
            t.urgency.title(),
            t.ticket_type.name  if t.ticket_type  else "",
            t.business_unit.name if t.business_unit else "",
            t.use_case.name     if t.use_case     else "",
            t.submitted_by.name if t.submitted_by else "",
            t.submitted_at.strftime("%Y-%m-%d %H:%M") if t.submitted_at else "",
            t.workstream_lead.name if t.workstream_lead else "",
            t.team_lead.name       if t.team_lead       else "",
            t.currently_with.name  if t.currently_with  else "",
            t.sla_deadline.strftime("%Y-%m-%d %H:%M") if t.sla_deadline else "",
            t.completed_at.strftime("%Y-%m-%d %H:%M") if t.completed_at else "",
        ]

    if format == "excel":
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "LiveOps Tickets"

        header_fill = PatternFill("solid", fgColor="1E40AF")
        header_font = Font(color="FFFFFF", bold=True, size=10)

        ws.append(headers_row)
        for cell in ws[1]:
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = Alignment(horizontal="center")

        for t in tickets:
            ws.append(_row(t))

        col_widths = [10, 16, 10, 24, 14, 24, 18, 18, 18, 18, 18, 18, 18]
        for i, width in enumerate(col_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"liveops_{from_date}_to_{to_date}.xlsx"
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    elif format == "pdf":
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import cm

        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1*cm, rightMargin=1*cm, topMargin=1.5*cm, bottomMargin=1.5*cm)
        styles = getSampleStyleSheet()

        title = Paragraph(f"LiveOps Report — {from_date} to {to_date}", styles["Title"])
        subtitle = Paragraph(f"Generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M')} UTC · {len(tickets)} tickets", styles["Normal"])

        pdf_headers = ["#", "Status", "Urgency", "Type", "BU", "Use Case", "Submitted By", "Submitted At", "WL", "PL", "Assignee", "SLA", "Completed"]
        data = [pdf_headers] + [_row(t) for t in tickets]

        col_widths_pdf = [1.2*cm, 2.2*cm, 1.8*cm, 3.5*cm, 1.8*cm, 3*cm, 2.5*cm, 2.5*cm, 2.5*cm, 2.5*cm, 2.5*cm, 2.5*cm, 2.5*cm]
        table = Table(data, colWidths=col_widths_pdf, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E40AF")),
            ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
            ("FONTSIZE",   (0, 0), (-1, 0), 7),
            ("FONTSIZE",   (0, 1), (-1, -1), 6.5),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
            ("GRID",       (0, 0), (-1, -1), 0.3, colors.HexColor("#E2E8F0")),
            ("VALIGN",     (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))

        doc.build([title, Spacer(1, 0.3*cm), subtitle, Spacer(1, 0.5*cm), table])
        buf.seek(0)
        filename = f"liveops_{from_date}_to_{to_date}.pdf"
        return StreamingResponse(
            buf,
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    raise HTTPException(status_code=422, detail="format must be 'excel' or 'pdf'")


# ── analytics ─────────────────────────────────────────────────────────────────

@router.get("/analytics/summary", response_model=LiveOpsAnalyticsSummary)
def analytics_summary(db: Session = Depends(get_db)):
    from sqlalchemy import func as sqlfunc
    now = _now()
    month_start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)

    total_open = db.query(LiveOpsTicket).filter(
        LiveOpsTicket.status.in_(["open", "in_progress"])
    ).count()

    pending_approval = db.query(LiveOpsTicket).filter(
        LiveOpsTicket.status.in_(["pending_wl", "pending_tl", "pending_lo"])
    ).count()

    in_progress = db.query(LiveOpsTicket).filter_by(status="in_progress").count()

    sla_breaching = db.query(LiveOpsTicket).filter(
        LiveOpsTicket.sla_deadline.isnot(None),
        LiveOpsTicket.sla_deadline < now,
        LiveOpsTicket.status.notin_(["completed", "rejected", "cancelled"]),
    ).count()

    completed_this_month = db.query(LiveOpsTicket).filter(
        LiveOpsTicket.status == "completed",
        LiveOpsTicket.completed_at >= month_start,
    ).count()

    completed = db.query(LiveOpsTicket).filter(
        LiveOpsTicket.status == "completed",
        LiveOpsTicket.completed_at >= month_start,
        LiveOpsTicket.submitted_at.isnot(None),
    ).all()

    avg_hours = None
    if completed:
        total_seconds = sum(
            (t.completed_at - t.submitted_at).total_seconds()
            for t in completed
            if t.completed_at and t.submitted_at
        )
        avg_hours = round(total_seconds / len(completed) / 3600, 1)

    return LiveOpsAnalyticsSummary(
        total_open=total_open,
        pending_approval=pending_approval,
        in_progress=in_progress,
        sla_breaching=sla_breaching,
        completed_this_month=completed_this_month,
        avg_resolution_hours=avg_hours,
    )
