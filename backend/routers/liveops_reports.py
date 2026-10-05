"""
LiveOps report generation — Excel and PDF export for PMO calls.
"""
import io
from datetime import datetime, timezone, date
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload

from backend.db.session import get_db
from backend.db.models import LiveOpsTicket, LiveOpsAssignment

router = APIRouter(prefix="/api/v1/liveops/reports", tags=["liveops-reports"])

# ── colour palette ─────────────────────────────────────────────────────────────

# Excel hex colours (no leading #)
XL_BLUE_DARK  = "1E3A5F"   # header bg
XL_BLUE_MID   = "2563EB"   # accent
XL_BLUE_LIGHT = "DBEAFE"   # alt row
XL_WHITE      = "FFFFFF"
XL_SLATE_50   = "F8FAFC"
XL_SLATE_200  = "E2E8F0"
XL_SLATE_700  = "334155"

STATUS_XL: dict[str, tuple[str, str]] = {
    "draft":       ("F8FAFC", "64748B"),
    "pending_pl":  ("FFFBEB", "B45309"),
    "pending_dl":  ("FFFBEB", "B45309"),
    "pending_lo":  ("FFF7ED", "C2410C"),
    "open":        ("EFF6FF", "1D4ED8"),
    "in_progress": ("EFF6FF", "2563EB"),
    "completed":   ("F0FDF4", "15803D"),
    "rejected":    ("FEF2F2", "B91C1C"),
    "cancelled":   ("F8FAFC", "64748B"),
    "on_hold":     ("FFFBEB", "B45309"),
}

URGENCY_XL: dict[str, tuple[str, str]] = {
    "high":   ("FEF2F2", "B91C1C"),
    "medium": ("FFFBEB", "B45309"),
    "low":    ("F0FDF4", "15803D"),
}

STATUS_LABEL = {
    "draft": "Draft", "pending_pl": "Pending PL", "pending_dl": "Pending DL",
    "pending_lo": "Pending LO", "open": "Open", "in_progress": "In Progress",
    "completed": "Completed", "rejected": "Rejected",
    "cancelled": "Cancelled", "on_hold": "On Hold",
}

# ReportLab colours
RL_BLUE_DARK  = (0.118, 0.227, 0.373)   # #1E3A5F
RL_BLUE_MID   = (0.145, 0.388, 0.922)   # #2563EB
RL_BLUE_LIGHT = (0.859, 0.918, 0.996)   # #DBEAFE
RL_WHITE      = (1, 1, 1)
RL_SLATE_50   = (0.973, 0.980, 0.988)
RL_SLATE_200  = (0.886, 0.910, 0.941)

STATUS_RL: dict[str, tuple] = {
    "draft":       ((0.973, 0.980, 0.988), (0.392, 0.455, 0.545)),
    "pending_pl":  ((1.0,   0.984, 0.918), (0.706, 0.349, 0.035)),
    "pending_dl":  ((1.0,   0.984, 0.918), (0.706, 0.349, 0.035)),
    "pending_lo":  ((1.0,   0.969, 0.929), (0.761, 0.255, 0.047)),
    "open":        ((0.937, 0.965, 1.0),   (0.114, 0.306, 0.871)),
    "in_progress": ((0.937, 0.965, 1.0),   (0.145, 0.388, 0.922)),
    "completed":   ((0.941, 0.992, 0.953), (0.086, 0.502, 0.239)),
    "rejected":    ((1.0,   0.949, 0.949), (0.725, 0.110, 0.110)),
    "cancelled":   ((0.973, 0.980, 0.988), (0.392, 0.455, 0.545)),
    "on_hold":     ((1.0,   0.984, 0.918), (0.706, 0.349, 0.035)),
}

URGENCY_RL: dict[str, tuple] = {
    "high":   ((1.0,   0.949, 0.949), (0.725, 0.110, 0.110)),
    "medium": ((1.0,   0.984, 0.918), (0.706, 0.349, 0.035)),
    "low":    ((0.941, 0.992, 0.953), (0.086, 0.502, 0.239)),
}


# ── helpers ────────────────────────────────────────────────────────────────────

def _parse_date(s: str) -> date:
    try:
        return datetime.strptime(s, "%Y-%m-%d").date()
    except ValueError:
        raise HTTPException(status_code=422, detail=f"Invalid date format: {s}. Use YYYY-MM-DD.")


def _fetch_tickets(db: Session, from_dt: datetime, to_dt: datetime):
    return (
        db.query(LiveOpsTicket)
        .options(
            joinedload(LiveOpsTicket.ticket_type),
            joinedload(LiveOpsTicket.use_case),
            joinedload(LiveOpsTicket.business_unit),
            joinedload(LiveOpsTicket.submitted_by),
            joinedload(LiveOpsTicket.practice_lead),
            joinedload(LiveOpsTicket.delivery_lead),
            joinedload(LiveOpsTicket.currently_with),
            joinedload(LiveOpsTicket.assignments).joinedload(LiveOpsAssignment.assignee),
        )
        .filter(LiveOpsTicket.created_at >= from_dt, LiveOpsTicket.created_at <= to_dt)
        .order_by(LiveOpsTicket.ticket_number)
        .all()
    )


def _duration_hours(ticket) -> str:
    if ticket.submitted_at and ticket.completed_at:
        delta = ticket.completed_at - ticket.submitted_at
        h = round(delta.total_seconds() / 3600, 1)
        return f"{h}h"
    return "—"


def _assignees(ticket) -> str:
    active = [a for a in ticket.assignments if a.status in ("active", "completed")]
    if not active:
        pending = [a for a in ticket.assignments if a.status == "pending"]
        names = [a.assignee.name for a in pending if a.assignee]
        return ", ".join(names) if names else "—"
    names = [a.assignee.name for a in active if a.assignee]
    return ", ".join(names) if names else "—"


def _fmt_dt(dt) -> str:
    if not dt:
        return "—"
    if hasattr(dt, "strftime"):
        return dt.strftime("%d %b %Y %H:%M")
    return str(dt)


def _fmt_date(dt) -> str:
    if not dt:
        return "—"
    if hasattr(dt, "strftime"):
        return dt.strftime("%d %b %Y")
    return str(dt)


def _build_stats(tickets: list) -> dict:
    total = len(tickets)
    by_status: dict[str, int] = {}
    by_urgency: dict[str, int] = {}
    by_bu: dict[str, int] = {}
    completed = [t for t in tickets if t.status == "completed"]
    breached = 0
    durations = []

    now = datetime.now(timezone.utc)
    for t in tickets:
        by_status[t.status] = by_status.get(t.status, 0) + 1
        by_urgency[t.urgency] = by_urgency.get(t.urgency, 0) + 1
        bu_name = t.business_unit.name if t.business_unit else "Unknown"
        by_bu[bu_name] = by_bu.get(bu_name, 0) + 1
        if t.sla_deadline and t.status not in ("completed", "rejected", "cancelled"):
            if t.sla_deadline < now:
                breached += 1
        if t.submitted_at and t.completed_at:
            delta = t.completed_at - t.submitted_at
            durations.append(delta.total_seconds() / 3600)

    avg_hours = round(sum(durations) / len(durations), 1) if durations else None
    sla_rate = round((1 - breached / total) * 100, 1) if total else 100.0

    return {
        "total": total,
        "completed": len(completed),
        "breached": breached,
        "avg_hours": avg_hours,
        "sla_rate": sla_rate,
        "by_status": by_status,
        "by_urgency": by_urgency,
        "by_bu": by_bu,
    }


# ── Excel ──────────────────────────────────────────────────────────────────────

def _build_excel(tickets: list, stats: dict, from_date: date, to_date: date) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import (
        PatternFill, Font, Alignment, Border, Side, numbers
    )
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # ── Sheet 1: Summary ──────────────────────────────────────────────────────

    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_sum.sheet_view.showGridLines = False
    ws_sum.column_dimensions["A"].width = 30
    ws_sum.column_dimensions["B"].width = 20
    ws_sum.column_dimensions["C"].width = 20

    hdr_fill   = PatternFill("solid", fgColor=XL_BLUE_DARK)
    hdr_font   = Font(color=XL_WHITE, bold=True, size=11, name="Calibri")
    title_font = Font(color=XL_BLUE_DARK, bold=True, size=16, name="Calibri")
    sub_font   = Font(color=XL_SLATE_700, size=10, name="Calibri")
    section_font = Font(color=XL_BLUE_DARK, bold=True, size=11, name="Calibri")
    cell_font  = Font(color=XL_SLATE_700, size=10, name="Calibri")
    thin_border = Border(
        bottom=Side(style="thin", color=XL_SLATE_200)
    )

    def hdr_cell(ws, row, col, value):
        c = ws.cell(row=row, column=col, value=value)
        c.fill = hdr_fill; c.font = hdr_font
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    def section_cell(ws, row, value):
        c = ws.cell(row=row, column=1, value=value)
        c.font = section_font
        ws.row_dimensions[row].height = 20

    def kv_row(ws, row, key, value, alt=False):
        k = ws.cell(row=row, column=1, value=key)
        v = ws.cell(row=row, column=2, value=value)
        k.font = cell_font; v.font = Font(color=XL_SLATE_700, size=10, bold=True, name="Calibri")
        k.alignment = Alignment(indent=2)
        v.alignment = Alignment(horizontal="right")
        if alt:
            fill = PatternFill("solid", fgColor=XL_BLUE_LIGHT)
            k.fill = fill; v.fill = fill
        k.border = thin_border; v.border = thin_border

    # Title
    ws_sum["A1"] = "LiveOps Ticketing — PMO Summary Report"
    ws_sum["A1"].font = title_font
    ws_sum["A1"].alignment = Alignment(vertical="center")
    ws_sum.row_dimensions[1].height = 32

    ws_sum["A2"] = f"Period: {from_date.strftime('%d %b %Y')} – {to_date.strftime('%d %b %Y')}"
    ws_sum["A2"].font = sub_font
    ws_sum["B2"] = f"Generated: {datetime.now().strftime('%d %b %Y %H:%M')}"
    ws_sum["B2"].font = sub_font
    ws_sum["B2"].alignment = Alignment(horizontal="right")
    ws_sum.row_dimensions[2].height = 18

    ws_sum.row_dimensions[3].height = 10

    # Overview section
    section_cell(ws_sum, 4, "Overview")
    ws_sum.row_dimensions[4].height = 22
    kv_row(ws_sum, 5, "Total Tickets",          stats["total"],     False)
    kv_row(ws_sum, 6, "Completed",               stats["completed"], True)
    kv_row(ws_sum, 7, "SLA Compliance Rate",     f"{stats['sla_rate']}%", False)
    kv_row(ws_sum, 8, "Active SLA Breaches",     stats["breached"],  True)
    kv_row(ws_sum, 9, "Avg Resolution Time",
           f"{stats['avg_hours']}h" if stats["avg_hours"] else "—", False)

    ws_sum.row_dimensions[10].height = 10

    # By Status section
    section_cell(ws_sum, 11, "By Status")
    ws_sum.row_dimensions[11].height = 22
    for i, (status, count) in enumerate(sorted(stats["by_status"].items())):
        kv_row(ws_sum, 12 + i, STATUS_LABEL.get(status, status), count, i % 2 == 1)
    row_after_status = 12 + len(stats["by_status"]) + 1

    section_cell(ws_sum, row_after_status, "By Urgency")
    ws_sum.row_dimensions[row_after_status].height = 22
    for i, (urg, count) in enumerate(sorted(stats["by_urgency"].items())):
        kv_row(ws_sum, row_after_status + 1 + i, urg.capitalize(), count, i % 2 == 1)
    row_after_urg = row_after_status + len(stats["by_urgency"]) + 2

    section_cell(ws_sum, row_after_urg, "By Business Unit")
    ws_sum.row_dimensions[row_after_urg].height = 22
    for i, (bu, count) in enumerate(sorted(stats["by_bu"].items())):
        kv_row(ws_sum, row_after_urg + 1 + i, bu, count, i % 2 == 1)

    # ── Sheet 2: Tickets ──────────────────────────────────────────────────────

    ws_tick = wb.create_sheet("Tickets")
    ws_tick.sheet_view.showGridLines = False
    ws_tick.freeze_panes = "A2"

    COLUMNS = [
        ("Ticket #",       12),
        ("Type",           22),
        ("Business Unit",  18),
        ("Use Case",       22),
        ("Urgency",        12),
        ("Status",         15),
        ("Submitted By",   20),
        ("Submitted",      18),
        ("SLA Deadline",   18),
        ("Practice Lead",  20),
        ("Delivery Lead",  20),
        ("Assignee(s)",    22),
        ("Resolved",       18),
        ("Duration",       12),
    ]

    for col, (title, width) in enumerate(COLUMNS, start=1):
        ws_tick.column_dimensions[get_column_letter(col)].width = width
        hdr_cell(ws_tick, 1, col, title)
    ws_tick.row_dimensions[1].height = 22

    for row_idx, t in enumerate(tickets, start=2):
        alt = row_idx % 2 == 0
        base_fill = PatternFill("solid", fgColor=XL_BLUE_LIGHT if alt else XL_WHITE)
        base_font = Font(color=XL_SLATE_700, size=9, name="Calibri")
        base_align = Alignment(vertical="center", wrap_text=False, indent=1)

        def put(col, value, fill=None, font=None, align=None):
            c = ws_tick.cell(row=row_idx, column=col, value=value)
            c.fill   = fill  or base_fill
            c.font   = font  or base_font
            c.alignment = align or base_align
            c.border = Border(bottom=Side(style="hair", color=XL_SLATE_200))

        put(1, f"LO-{t.ticket_number:04d}")
        put(2, t.ticket_type.name if t.ticket_type else "—")
        put(3, t.business_unit.name if t.business_unit else "—")
        put(4, t.use_case.name if t.use_case else "—")

        # Urgency cell
        urg_bg, urg_fg = URGENCY_XL.get(t.urgency, (XL_WHITE, XL_SLATE_700))
        put(5, t.urgency.capitalize(),
            fill=PatternFill("solid", fgColor=urg_bg),
            font=Font(color=urg_fg, size=9, bold=True, name="Calibri"),
            align=Alignment(horizontal="center", vertical="center"))

        # Status cell
        st_bg, st_fg = STATUS_XL.get(t.status, (XL_WHITE, XL_SLATE_700))
        put(6, STATUS_LABEL.get(t.status, t.status),
            fill=PatternFill("solid", fgColor=st_bg),
            font=Font(color=st_fg, size=9, bold=True, name="Calibri"),
            align=Alignment(horizontal="center", vertical="center"))

        put(7, t.submitted_by.name if t.submitted_by else "—")
        put(8, _fmt_date(t.submitted_at))
        put(9, _fmt_date(t.sla_deadline))
        put(10, t.practice_lead.name if t.practice_lead else "—")
        put(11, t.delivery_lead.name if t.delivery_lead else "—")
        put(12, _assignees(t))
        put(13, _fmt_date(t.completed_at))
        put(14, _duration_hours(t))

        ws_tick.row_dimensions[row_idx].height = 16

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ── PDF ────────────────────────────────────────────────────────────────────────

def _build_pdf(tickets: list, stats: dict, from_date: date, to_date: date) -> bytes:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
    )
    from reportlab.graphics.shapes import Drawing, Rect
    from reportlab.graphics.charts.barcharts import VerticalBarChart

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=1.5*cm, rightMargin=1.5*cm,
        topMargin=1.5*cm, bottomMargin=1.5*cm,
    )

    C_BLUE_DARK  = colors.HexColor("#1E3A5F")
    C_BLUE_MID   = colors.HexColor("#2563EB")
    C_BLUE_LIGHT = colors.HexColor("#DBEAFE")
    C_WHITE      = colors.white
    C_SLATE_50   = colors.HexColor("#F8FAFC")
    C_SLATE_200  = colors.HexColor("#E2E8F0")
    C_SLATE_600  = colors.HexColor("#475569")
    C_SLATE_800  = colors.HexColor("#1E293B")
    C_GREEN      = colors.HexColor("#15803D")
    C_RED        = colors.HexColor("#B91C1C")
    C_AMBER      = colors.HexColor("#B45309")

    STATUS_PDF_BG  = {k: colors.HexColor(f"#{v[0]}") for k, v in STATUS_XL.items()}
    STATUS_PDF_FG  = {k: colors.HexColor(f"#{v[1]}") for k, v in STATUS_XL.items()}
    URGENCY_PDF_BG = {k: colors.HexColor(f"#{v[0]}") for k, v in URGENCY_XL.items()}
    URGENCY_PDF_FG = {k: colors.HexColor(f"#{v[1]}") for k, v in URGENCY_XL.items()}

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=18, textColor=C_BLUE_DARK, spaceAfter=4)
    sub_style   = ParagraphStyle("sub",   fontName="Helvetica",      fontSize=9,  textColor=C_SLATE_600)
    section_style = ParagraphStyle("section", fontName="Helvetica-Bold", fontSize=11, textColor=C_BLUE_DARK, spaceBefore=12, spaceAfter=6)
    cell_style  = ParagraphStyle("cell",  fontName="Helvetica",      fontSize=7.5, textColor=C_SLATE_800, leading=10)

    story = []

    # ── Title block ───────────────────────────────────────────────────────────
    story.append(Paragraph("LiveOps Ticketing — PMO Summary Report", title_style))
    story.append(Paragraph(
        f"Period: <b>{from_date.strftime('%d %b %Y')}</b> – <b>{to_date.strftime('%d %b %Y')}</b> &nbsp;&nbsp; "
        f"Generated: {datetime.now().strftime('%d %b %Y %H:%M')}",
        sub_style
    ))
    story.append(HRFlowable(width="100%", thickness=1.5, color=C_BLUE_MID, spaceAfter=10))

    # ── KPI row ───────────────────────────────────────────────────────────────
    kpi_data = [
        ["Total Tickets", "Completed", "SLA Compliance", "Active SLA Breaches", "Avg Resolution"],
        [
            str(stats["total"]),
            str(stats["completed"]),
            f"{stats['sla_rate']}%",
            str(stats["breached"]),
            f"{stats['avg_hours']}h" if stats["avg_hours"] else "—",
        ],
    ]
    kpi_table = Table(kpi_data, colWidths=[5.6*cm]*5, rowHeights=[0.8*cm, 1.4*cm])
    kpi_table.setStyle(TableStyle([
        ("BACKGROUND",   (0,0), (-1,0), C_BLUE_DARK),
        ("TEXTCOLOR",    (0,0), (-1,0), C_WHITE),
        ("FONTNAME",     (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTSIZE",     (0,0), (-1,0), 8),
        ("BACKGROUND",   (0,1), (-1,1), C_SLATE_50),
        ("FONTNAME",     (0,1), (-1,1), "Helvetica-Bold"),
        ("FONTSIZE",     (0,1), (-1,1), 18),
        ("TEXTCOLOR",    (0,1), (-1,1), C_BLUE_DARK),
        ("ALIGN",        (0,0), (-1,-1), "CENTER"),
        ("VALIGN",       (0,0), (-1,-1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0,1), (-1,1), [C_SLATE_50]),
        ("BOX",          (0,0), (-1,-1), 1, C_SLATE_200),
        ("INNERGRID",    (0,0), (-1,-1), 0.5, C_SLATE_200),
        ("ROUNDEDCORNERS", [4]),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 10))

    # ── Breakdown tables side by side ─────────────────────────────────────────
    def breakdown_table(title_text, data_dict, colour_map_bg=None, colour_map_fg=None):
        rows = [[title_text, ""]]
        for k, v in sorted(data_dict.items(), key=lambda x: -x[1]):
            label = STATUS_LABEL.get(k, k.capitalize())
            rows.append([label, str(v)])
        tbl = Table(rows, colWidths=[4.5*cm, 1.5*cm])
        style_cmds = [
            ("BACKGROUND",  (0,0), (-1,0), C_BLUE_DARK),
            ("TEXTCOLOR",   (0,0), (-1,0), C_WHITE),
            ("FONTNAME",    (0,0), (-1,0), "Helvetica-Bold"),
            ("FONTSIZE",    (0,0), (-1,-1), 8),
            ("ALIGN",       (1,1), (1,-1), "RIGHT"),
            ("VALIGN",      (0,0), (-1,-1), "MIDDLE"),
            ("BOX",         (0,0), (-1,-1), 1, C_SLATE_200),
            ("INNERGRID",   (0,0), (-1,-1), 0.5, C_SLATE_200),
            ("SPAN",        (0,0), (1,0)),
            ("ROWBACKGROUNDS", (0,1), (-1,-1), [C_WHITE, C_SLATE_50]),
        ]
        if colour_map_bg and colour_map_fg:
            for i, (k, _) in enumerate(sorted(data_dict.items(), key=lambda x: -x[1]), start=1):
                bg = colour_map_bg.get(k)
                fg = colour_map_fg.get(k)
                if bg:
                    style_cmds.append(("BACKGROUND", (0,i), (0,i), bg))
                if fg:
                    style_cmds.append(("TEXTCOLOR", (0,i), (0,i), fg))
        tbl.setStyle(TableStyle(style_cmds))
        return tbl

    tbl_status  = breakdown_table("By Status",  stats["by_status"],  STATUS_PDF_BG,  STATUS_PDF_FG)
    tbl_urgency = breakdown_table("By Urgency", stats["by_urgency"], URGENCY_PDF_BG, URGENCY_PDF_FG)
    tbl_bu      = breakdown_table("By Business Unit", stats["by_bu"])

    breakdown_row = Table([[tbl_status, tbl_urgency, tbl_bu]], colWidths=[6.5*cm, 6.5*cm, 8.5*cm])
    breakdown_row.setStyle(TableStyle([
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("LEFTPADDING",  (0,0), (-1,-1), 6),
        ("RIGHTPADDING", (0,0), (-1,-1), 6),
    ]))
    story.append(breakdown_row)
    story.append(Spacer(1, 14))

    # ── Ticket table ──────────────────────────────────────────────────────────
    story.append(HRFlowable(width="100%", thickness=0.5, color=C_SLATE_200, spaceBefore=2, spaceAfter=8))
    story.append(Paragraph("Ticket Listing", section_style))

    col_widths = [1.6*cm, 3.8*cm, 2.8*cm, 3.8*cm, 1.8*cm, 2.2*cm, 3.2*cm, 2.4*cm, 2.4*cm, 3.2*cm, 1.8*cm]

    headers = ["Ticket #", "Type", "Business Unit", "Use Case", "Urgency", "Status",
               "Submitted By", "Submitted", "SLA Deadline", "Assignee(s)", "Duration"]

    table_data = [headers]
    for t in tickets:
        urg_label  = t.urgency.capitalize()
        stat_label = STATUS_LABEL.get(t.status, t.status)
        table_data.append([
            f"LO-{t.ticket_number:04d}",
            t.ticket_type.name if t.ticket_type else "—",
            t.business_unit.name if t.business_unit else "—",
            t.use_case.name if t.use_case else "—",
            urg_label,
            stat_label,
            t.submitted_by.name if t.submitted_by else "—",
            _fmt_date(t.submitted_at),
            _fmt_date(t.sla_deadline),
            _assignees(t),
            _duration_hours(t),
        ])

    ticket_table = Table(table_data, colWidths=col_widths, repeatRows=1)

    ts = TableStyle([
        # Header
        ("BACKGROUND",  (0,0), (-1,0), C_BLUE_DARK),
        ("TEXTCOLOR",   (0,0), (-1,0), C_WHITE),
        ("FONTNAME",    (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTSIZE",    (0,0), (-1,0), 7.5),
        ("ROWHEIGHT",   0, 16),
        # Body
        ("FONTNAME",    (0,1), (-1,-1), "Helvetica"),
        ("FONTSIZE",    (0,1), (-1,-1), 7),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [C_WHITE, C_SLATE_50]),
        ("VALIGN",      (0,0), (-1,-1), "MIDDLE"),
        ("ALIGN",       (0,0), (-1,-1), "LEFT"),
        ("ALIGN",       (10,0), (10,-1), "CENTER"),
        ("BOX",         (0,0), (-1,-1), 1, C_SLATE_200),
        ("INNERGRID",   (0,0), (-1,-1), 0.25, C_SLATE_200),
        ("LEFTPADDING",  (0,0), (-1,-1), 4),
        ("RIGHTPADDING", (0,0), (-1,-1), 4),
        ("TOPPADDING",   (0,0), (-1,-1), 2),
        ("BOTTOMPADDING",(0,0), (-1,-1), 2),
    ])

    # Color urgency + status cells per row
    for row_idx, t in enumerate(tickets, start=1):
        urg_bg = URGENCY_PDF_BG.get(t.urgency)
        urg_fg = URGENCY_PDF_FG.get(t.urgency)
        st_bg  = STATUS_PDF_BG.get(t.status)
        st_fg  = STATUS_PDF_FG.get(t.status)
        if urg_bg:
            ts.add("BACKGROUND", (4, row_idx), (4, row_idx), urg_bg)
            ts.add("TEXTCOLOR",  (4, row_idx), (4, row_idx), urg_fg or C_SLATE_800)
            ts.add("FONTNAME",   (4, row_idx), (4, row_idx), "Helvetica-Bold")
        if st_bg:
            ts.add("BACKGROUND", (5, row_idx), (5, row_idx), st_bg)
            ts.add("TEXTCOLOR",  (5, row_idx), (5, row_idx), st_fg or C_SLATE_800)
            ts.add("FONTNAME",   (5, row_idx), (5, row_idx), "Helvetica-Bold")

    ticket_table.setStyle(ts)
    story.append(ticket_table)

    doc.build(story)
    buf.seek(0)
    return buf.read()


# ── endpoint ──────────────────────────────────────────────────────────────────

@router.get("/export")
def export_report(
    from_date: str = Query(..., description="Start date YYYY-MM-DD"),
    to_date:   str = Query(..., description="End date YYYY-MM-DD (inclusive)"),
    format:    str = Query("excel", description="'excel' or 'pdf'"),
    db: Session = Depends(get_db),
):
    """Generate a PMO summary report for the given period."""
    fmt = format.lower()
    if fmt not in ("excel", "pdf"):
        raise HTTPException(status_code=422, detail="format must be 'excel' or 'pdf'")

    fd = _parse_date(from_date)
    td = _parse_date(to_date)
    if td < fd:
        raise HTTPException(status_code=422, detail="to_date must be >= from_date")

    from_dt = datetime(fd.year, fd.month, fd.day, 0, 0, 0, tzinfo=timezone.utc)
    to_dt   = datetime(td.year, td.month, td.day, 23, 59, 59, tzinfo=timezone.utc)

    tickets = _fetch_tickets(db, from_dt, to_dt)
    stats   = _build_stats(tickets)

    period_label = f"{fd.strftime('%Y-%m-%d')}_to_{td.strftime('%Y-%m-%d')}"

    if fmt == "excel":
        data = _build_excel(tickets, stats, fd, td)
        filename = f"liveops_report_{period_label}.xlsx"
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        data = _build_pdf(tickets, stats, fd, td)
        filename = f"liveops_report_{period_label}.pdf"
        media_type = "application/pdf"

    return StreamingResponse(
        io.BytesIO(data),
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
